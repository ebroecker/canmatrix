import copy
import sys
from decimal import Decimal

import openpyxl
import pytest
from click.testing import CliRunner

import canmatrix.compare
import canmatrix.formats
from canmatrix._canmatrix import CanMatrix
from canmatrix.Frame import Frame
from canmatrix.Signal import Signal
from canmatrix.cli.compare import cli_compare


def comparison_matrices():
    old = CanMatrix()
    old.add_frame_defines("Variant", "STRING")
    frame = Frame("Engine", arbitration_id=0x123, size=8,
                  transmitters=["OldECU"], comment="Old comment")
    frame.add_signal(Signal("Speed", start_bit=0, size=8, factor=Decimal("0.1"),
                            receivers=["OldReceiver"], values={1: "One"}))
    frame.add_signal(Signal("Removed", start_bit=8, size=8))
    frame.add_attribute("Variant", "Old")
    old.add_frame(frame)
    new = copy.deepcopy(old)
    frame = new.frames[0]
    frame.transmitters = ["NewECU"]
    frame.comment = "New comment"
    frame.attributes["Variant"] = "New"
    frame.signals[0].factor = Decimal("0.2")
    frame.signals[0].receivers = ["NewReceiver"]
    frame.signals[0].values = {1: "Two"}
    frame.signals.pop()
    frame.add_signal(Signal("Added", start_bit=8, size=8))
    return old, new


def report_rows(path):
    book = openpyxl.load_workbook(path)
    try:
        return list(book.active.iter_rows(values_only=True))
    finally:
        book.close()


def test_xlsx_report_values_context_and_no_mutation(tmp_path):
    old, new = comparison_matrices()
    result = canmatrix.compare.compare_db(old, new)
    before = repr(result)
    output = tmp_path / "changes.xlsx"
    canmatrix.compare.dump_result_xlsx(result, output)
    rows = report_rows(output)
    assert rows[0] == ("Object", "Property", "Change", "Old value", "New value")
    assert ("FRAME Engine / SIGNAL Speed", "factor", "changed", "0.1", "0.2") in rows
    assert ("FRAME Engine", "Frame-Transmitter", "removed", "OldECU", None) in rows
    assert ("FRAME Engine", "Frame-Transmitter", "added", None, "NewECU") in rows
    assert ("FRAME Engine / SIGNAL Speed", "receiver OldReceiver", "removed",
            "OldReceiver", None) in rows
    assert ("FRAME Engine / SIGNAL Speed", "receiver NewReceiver", "added",
            None, "NewReceiver") in rows
    assert any(row[2:] == ("changed", "One", "Two") for row in rows[1:])
    assert any(row[1:4] == ("SIGNAL", "added", None) and row[4] == "Added"
               for row in rows[1:])
    assert ("FRAME Engine / SIGNAL Removed", "SIGNAL", "deleted", "Removed", None) in rows
    assert repr(result) == before
    canmatrix.compare.dump_result_xlsx(result, tmp_path / "again.xlsx")
    assert report_rows(tmp_path / "again.xlsx") == rows


def test_xlsx_report_preserves_literal_unicode_text(tmp_path):
    result = canmatrix.compare.CompareResult("changed", "comment", None,
                                            ["=1+1", "Größe 日本語"])
    output = tmp_path / "text.xlsx"
    canmatrix.compare.dump_result_xlsx(result, output)
    book = openpyxl.load_workbook(output)
    try:
        assert book.active["D2"].value == "=1+1"
        assert book.active["D2"].data_type == "s"
        assert book.active["E2"].value == "Größe 日本語"
    finally:
        book.close()


def test_xlsx_report_equal_matrices(tmp_path):
    matrix = CanMatrix()
    output = tmp_path / "equal.xlsx"
    canmatrix.compare.dump_result_xlsx(canmatrix.compare.compare_db(matrix, matrix), output)
    assert report_rows(output) == [("Object", "Property", "Change", "Old value", "New value")]


def test_xlsx_report_definition_and_table_values(tmp_path):
    old, new = CanMatrix(), CanMatrix()
    old.add_frame_defines("OldDefinition", "INT 0 7")
    new.add_frame_defines("NewDefinition", "STRING")
    old.value_tables["OldTable"] = {1: "One"}
    new.value_tables["NewTable"] = {2: "Two"}
    output = tmp_path / "definitions.xlsx"
    canmatrix.compare.dump_result_xlsx(canmatrix.compare.compare_db(old, new), output)
    rows = report_rows(output)
    assert ("Frame Defines", "DefineOldDefinition", "deleted", "INT 0 7", None) in rows
    assert ("Frame Defines", "DefineNewDefinition", "added", None, "STRING") in rows
    assert ("Database", "valuetable OldTable", "deleted", "{1: 'One'}", None) in rows
    assert ("Database", "valuetable NewTable", "added", None, "{2: 'Two'}") in rows


def test_xlsx_report_preserves_precision_and_cell_limit(tmp_path):
    result = canmatrix.compare.CompareResult()
    result.add_child(canmatrix.compare.CompareResult("changed", "offset", None,
                     [Decimal("0.123456789012345678901234567890"), 9007199254740993]))
    result.add_child(canmatrix.compare.CompareResult("changed", "comment", None,
                     ["old", "x" * 32767]))
    output = tmp_path / "precision.xlsx"
    canmatrix.compare.dump_result_xlsx(result, output)
    rows = report_rows(output)
    assert rows[1][3:] == ("0.123456789012345678901234567890", "9007199254740993")
    assert rows[2][4] == "x" * 32767


def test_xlsx_report_rejects_oversize_text_before_saving(tmp_path):
    output = tmp_path / "existing.xlsx"
    output.write_bytes(b"previous report")
    result = canmatrix.compare.CompareResult("changed", "comment", None, ["old", "x" * 32768])
    with pytest.raises(ValueError, match="32767"):
        canmatrix.compare.dump_result_xlsx(result, output)
    assert output.read_bytes() == b"previous report"


def test_cli_xlsx_report_respects_filters_and_keeps_console(tmp_path):
    old, new = comparison_matrices()
    inputs = [tmp_path / "old.json", tmp_path / "new.json"]
    for matrix, path in zip((old, new), inputs):
        canmatrix.formats.dumpp({"": matrix}, str(path), jsonExportAll=True)
    output = tmp_path / "changes.xlsx"
    runner = CliRunner()
    normal = runner.invoke(cli_compare, list(map(str, inputs)))
    exported = runner.invoke(cli_compare, ["--output", str(output)] + list(map(str, inputs)))
    assert normal.exit_code == exported.exit_code == 0
    assert normal.stdout == exported.stdout
    rows = report_rows(output)
    assert not any(row[1] == "Variant" for row in rows)
    assert not any(row[3] == "comment: Old comment" for row in rows)
    assert any(row[2:] == ("changed", "One", "Two") for row in rows)

    filtered = runner.invoke(cli_compare, ["--comments", "--attributes", "--valueTable",
                                          "--output", str(output)] + list(map(str, inputs)))
    assert filtered.exit_code == 0
    rows = report_rows(output)
    assert any(row[1:] == ("Variant", "changed", "Old", "New") for row in rows)
    assert any(row[3:] == ("comment: Old comment", "comment: New comment") for row in rows)
    assert not any(row[2:] == ("changed", "One", "Two") for row in rows)


@pytest.mark.parametrize("arguments", [["--frames"], []])
def test_cli_rejects_incompatible_report_options(tmp_path, arguments):
    output = tmp_path / ("changes.xlsx" if arguments else "changes.csv")
    result = CliRunner().invoke(cli_compare,
                                arguments + ["--output", str(output), "old.dbc", "new.dbc"])
    assert result.exit_code == 2
    expected = "--frames" if arguments else ".xlsx"
    assert expected in result.output
    assert not output.exists()


def test_cli_report_missing_optional_dependency(tmp_path, monkeypatch):
    monkeypatch.setitem(sys.modules, "openpyxl", None)
    output = tmp_path / "changes.xlsx"
    result = CliRunner().invoke(cli_compare, ["--output", str(output), "old.dbc", "new.dbc"])
    assert result.exit_code == 1
    assert "canmatrix[xlsx]" in result.output
    assert not output.exists()
